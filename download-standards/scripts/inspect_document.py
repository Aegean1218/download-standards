#!/usr/bin/env python3
"""Inspect a downloaded document; technical checks never prove currency or completeness."""
import argparse
import hashlib
import json
import os
import re
import unicodedata
import zipfile
from pathlib import Path


def fingerprint(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def report_conflicts(report, *paths):
    if report is None:
        return False
    report = Path(report).expanduser().resolve()
    for path in paths:
        if path is None:
            continue
        path = Path(path).expanduser().resolve()
        if report == path or (report.exists() and path.exists() and os.path.samefile(report, path)):
            return True
    return False


def document_format(path):
    with Path(path).open('rb') as stream:
        head = stream.read(1024)
    if head.startswith(b'%PDF-'):
        return 'pdf'
    if head.startswith(b'PK'):
        try:
            with zipfile.ZipFile(path) as archive:
                if {'[Content_Types].xml', 'word/document.xml'}.issubset(archive.namelist()):
                    return 'docx'
        except (OSError, zipfile.BadZipFile):
            pass
        return 'unknown'
    if head.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        return 'legacy_ole_unverified'
    return 'unknown'


def normalize(text):
    text = unicodedata.normalize('NFKC', text).upper()
    return re.sub(r'\s+', '', text.translate(str.maketrans({'—': '-', '–': '-', '−': '-'})))


def inspect_document(path, expected_standard=None, expected_pages=None, *, filename=None):
    path = Path(path)
    result = {'path': str(path.resolve()), 'status': 'downloaded_pending_review',
              'format': None, 'bytes': None, 'sha256': None, 'pages': None,
              'expected_standard': expected_standard, 'identity_match': None,
              'expected_pages': expected_pages, 'completeness': 'not_established',
              'watermark': 'not_visually_checked', 'warnings': []}
    try:
        result.update(bytes=path.stat().st_size, sha256=fingerprint(path), format=document_format(path))
        fmt = result['format']
        expected_suffix = {'pdf': '.pdf', 'docx': '.docx', 'legacy_ole_unverified': '.doc'}.get(fmt)
        label = Path(filename) if filename else path
        if expected_suffix and label.suffix.lower() not in (expected_suffix, '.part'):
            result['warnings'].append('extension_format_mismatch')
        text = ''
        if fmt == 'pdf':
            import fitz
            with fitz.open(path) as doc:
                if doc.needs_pass:
                    result['warnings'].append('password_or_access_restricted')
                    return result
                result['pages'] = len(doc)
                if not len(doc):
                    raise ValueError('empty_document')
                low_text = []
                cover = []
                for number, page in enumerate(doc, 1):
                    page_text = page.get_text()
                    if number <= 3:
                        cover.append(page_text)
                    if len(page_text.strip()) < 30:
                        low_text.append(number)
                text = '\n'.join(cover)
                result['low_text_pages'] = low_text
                if low_text:
                    result['warnings'].append('image_or_low_text_pages')
                result['repaired_by_parser'] = doc.is_repaired
                if doc.is_repaired:
                    result['warnings'].append('pdf_parser_repaired_file')
        elif fmt == 'docx':
            from docx import Document
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs[:80])
            text += '\n' + '\n'.join(cell.text for table in doc.tables[:3] for row in table.rows for cell in row.cells)
            result['warnings'].append('word_pagination_requires_visual_review')
        elif fmt == 'legacy_ole_unverified':
            result['warnings'].append('legacy_word_requires_native_app_review')
            return result
        else:
            raise ValueError('not_pdf_or_word')
        normalized = normalize(text)
        if expected_standard and normalized:
            # Prevent a shorter number/year from matching a longer identifier.
            pattern = re.escape(normalize(expected_standard)) + r'(?![0-9])'
            result['identity_match'] = bool(re.search(pattern, normalized))
            if not result['identity_match']:
                result['warnings'].append('standard_identity_mismatch')
        if expected_pages is not None and result['pages'] is not None and expected_pages != result['pages']:
            result['warnings'].append('page_count_mismatch')
        if re.search(r'征求意见稿|报批稿|送审稿|预览版|\bDRAFT\b(?:\s+FOR\s+COMMENT)?|\bPREVIEW\b', text, re.I):
            result['warnings'].append('draft_or_preview')
        if re.search(r'发布.*公告|标准目录|标准清单|标准批准发布公告', text):
            result['warnings'].append('possible_notice_or_catalog_check_body')
        if re.search(r'水印|仅供.*(?:阅读|学习)|禁止复制|WATERMARK', text, re.I):
            result['warnings'].append('possible_watermark_or_usage_notice')
    except ImportError as exc:
        result['warnings'].append('missing_dependency:' + (exc.name or 'unknown'))
    except Exception as exc:
        result['status'] = 'invalid_file'
        result['error'] = type(exc).__name__
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('path', type=Path)
    parser.add_argument('--expected-standard')
    parser.add_argument('--expected-pages', type=int)
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    if report_conflicts(args.report, args.path):
        parser.error('--report must not overwrite the input document')
    if args.expected_pages is not None and args.expected_pages < 1:
        parser.error('--expected-pages must be positive')
    result = inspect_document(args.path, args.expected_standard, args.expected_pages)
    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(output + '\n', encoding='utf-8')
    print(output)
    return 1 if result['status'] == 'invalid_file' else 0


if __name__ == '__main__':
    raise SystemExit(main())
